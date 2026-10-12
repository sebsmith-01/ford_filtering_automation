"""Weekly tagging-accuracy report.

Compares the analysts' work in a week's "For report" Google Sheet against the LLM tags in
that week's added_facebook_names file, writes an Excel report to accuracy_reports/ and posts
the headline metrics (with the Excel attached) to Slack.

Run from the repo root — defaults to last week's data, so it is scheduled for Monday mornings:
    python weekly_processing/tagging_accuracy.py
    python weekly_processing/tagging_accuracy.py --week 2026-10-05 --no-slack
"""
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import argparse
import datetime
import os

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from auth import get_service_account_creds
from google_sheet_processor import GoogleSheetProcessor, extract_file_id
from helper_functions import get_monday_str, get_spreadsheet_url

load_dotenv(PROJECT_ROOT / ".env")

SLACK_CHANNEL = "C0C7X7HG7SP"
REPORT_DIR = PROJECT_ROOT / "accuracy_reports"
HISTORY_PATH = REPORT_DIR / "accuracy_history.csv"

# Tagged irregularly, so not flagged when a week goes unreviewed
IRREGULAR_TABS = {"R4", "R5"}

# Until w/c 2026-10-12 some analysts marked only Hits and left rejected rows blank. From then on
# every tab should mark Misses too, so a tab with Hits but no Misses is flagged.
ALL_TABS_MARK_MISSES_FROM = "2026-10-12"

# LLM-tagged fields an analyst can correct on a Hit: (column, label used in the report)
TAG_FIELDS = [
    ("feedback_subcategory", "Subcategory"),
    ("feedback_sentiment", "Feedback sentiment"),
    ("overall_sentiment", "Overall sentiment"),
    ("ownership_status", "Ownership"),
    ("country_code", "Country"),
    ("is_malfunction", "Malfunction"),
    ("model_comparison", "Model comparison"),
]
BOOL_FIELDS = {"is_malfunction", "model_comparison"}

VALID_SENTIMENTS = {"Negative", "Neutral", "Positive"}
VALID_OWNERSHIP = {"Owner", "Pre-Ownership", "Showing Interest"}  # what excel_upload.py accepts on a Hit

# Small groups make noisy percentages — only rank subcategories with at least this many rows
MIN_HITS_TO_RANK = 10
MIN_REVIEWED_TO_RANK = 20

SHEET_COLS = ["id-3", "Validation", "url", "translated_text", "post_text", "Sidenote",
              "validation_auto", "confidence", "autoval_status", "is_malfunction_auto",
              "model_comparison_auto"] + [f for f, _ in TAG_FIELDS]


# ──────────────────────────────────────────────────────────────────────────────
# Loading
# ──────────────────────────────────────────────────────────────────────────────

def last_monday_str() -> str:
    this_monday = datetime.date.fromisoformat(get_monday_str())
    return (this_monday - datetime.timedelta(days=7)).isoformat()


def load_llm_tags(week: str) -> pd.DataFrame:
    """The LLM's tags for every row of the week, keyed by id-3 (unique per mention × subcategory)."""
    path = PROJECT_ROOT / f"weekly_data/{week}/added_facebook_names_{week}.xlsx"
    cols = ["id-3"] + [f for f, _ in TAG_FIELDS]
    llm = pd.read_excel(path, usecols=cols, dtype=str).fillna("")
    return llm.drop_duplicates("id-3").set_index("id-3")


def load_sheet(week: str) -> tuple[pd.DataFrame, list[str], list[str], str]:
    """Every vehicle tab's rows stacked into one frame, plus the tabs missing from or empty in the sheet."""
    url = get_spreadsheet_url(week)
    if url is None:
        raise FileNotFoundError(f"No sheet for w/c {week} in spreadsheet_links.csv")
    sheet = GoogleSheetProcessor(extract_file_id(url), get_service_account_creds())

    vehicle_tabs = pd.read_csv(PROJECT_ROOT / "vehicle_ids.csv")["vehicle_name"].tolist()
    missing = [t for t in vehicle_tabs if t not in sheet.tab_ids]
    empty, frames = [], []
    for tab in vehicle_tabs:
        if tab in missing:
            continue
        df = sheet.read_tab(tab)
        if df.empty or "id-3" not in df.columns:
            empty.append(tab)
            continue
        df = df.loc[:, ~df.columns.duplicated()]
        frames.append(df.reindex(columns=SHEET_COLS, fill_value="").assign(tab=tab))
    rows = pd.concat(frames, ignore_index=True).fillna("")
    return rows, missing, empty, url


# ──────────────────────────────────────────────────────────────────────────────
# Row-level preparation
# ──────────────────────────────────────────────────────────────────────────────

def normalise_label(s: pd.Series) -> pd.Series:
    """Hit / Miss / Maybe / "" (blank) / Invalid (anything else typed into the column)."""
    v = s.astype(str).str.strip().str.lower().str.replace(" ", "", regex=False)
    return pd.Series(
        np.select([v == "hit", v == "miss", v == "maybe", v == ""], ["Hit", "Miss", "Maybe", ""], "Invalid"),
        index=s.index,
    )


def normalise_tag(s: pd.Series, field: str) -> pd.Series:
    v = s.astype(str).str.strip().str.lower()
    if field in BOOL_FIELDS:
        return v.isin(["true", "1", "1.0", "yes"])
    return v


def untagged_fields(llm: pd.DataFrame) -> list[str]:
    """Fields holding a single value across the whole LLM file — the LLM didn't tag them that week
    (is_malfunction and model_comparison are always False upstream), so an analyst setting them is
    an addition, not a correction, and they're left out of tag accuracy."""
    return [f for f, _ in TAG_FIELDS if normalise_tag(llm[f], f).nunique() <= 1]


def prepare_rows(rows: pd.DataFrame, llm: pd.DataFrame, untagged: list[str]) -> pd.DataFrame:
    rows = rows.copy()
    rows["validation"] = normalise_label(rows["Validation"])
    rows["auto"] = normalise_label(rows["validation_auto"])
    rows["confidence"] = pd.to_numeric(rows["confidence"], errors="coerce")
    rows["in_llm_file"] = rows["id-3"].isin(llm.index)

    # Review mode per tab: a tab that marks Misses tells us what was reviewed; a Hits-only tab
    # (older weeks) doesn't, so its blanks can't be read as either Miss or unreviewed
    marks = rows.groupby("tab")["validation"].agg(lambda v: (v == "Miss").any() * 2 + (v == "Hit").any())
    mode = marks.map({3: "Hit+Miss", 2: "Hit+Miss", 1: "Hits only", 0: "Not reviewed"})
    rows["review_mode"] = rows["tab"].map(mode)
    rows["full_review"] = rows["review_mode"] == "Hit+Miss"

    rows["is_hit"] = rows["validation"] == "Hit"
    rows["is_miss"] = rows["validation"] == "Miss"
    rows["is_maybe"] = rows["validation"] == "Maybe"
    rows["reviewed"] = rows["is_hit"] | rows["is_miss"] | rows["is_maybe"]

    # The LLM's subcategory is what the row was tagged with, so group subcategory metrics by it
    llm_aligned = llm.reindex(rows["id-3"]).set_axis(rows.index)
    rows["llm_subcategory"] = llm_aligned["feedback_subcategory"].where(rows["in_llm_file"], rows["feedback_subcategory"])

    for field, _ in TAG_FIELDS:
        rows[f"llm_{field}"] = llm_aligned[field].fillna("")
        if field in untagged:
            rows[f"kept_{field}"] = np.nan
            continue
        changed = normalise_tag(rows[field], field) != normalise_tag(rows[f"llm_{field}"], field)
        # Only Hits are corrected; a row added by hand has nothing to compare against
        rows[f"kept_{field}"] = np.where(rows["is_hit"] & rows["in_llm_file"], ~changed, np.nan)
    kept_cols = [f"kept_{f}" for f, _ in TAG_FIELDS if f not in untagged]
    rows["kept_all"] = np.where(rows["is_hit"] & rows["in_llm_file"], rows[kept_cols].eq(1).all(axis=1), np.nan)
    return rows


# ──────────────────────────────────────────────────────────────────────────────
# Metrics
# ──────────────────────────────────────────────────────────────────────────────

def _rate(num, den):
    return num / den if den else np.nan


def summarise(g: pd.DataFrame) -> pd.Series:
    """Volume, relevance and tag-accuracy metrics for any group of rows."""
    full = g[g["full_review"]]
    hits, misses = int(full["is_hit"].sum()), int(full["is_miss"].sum())
    out = {
        "Rows": len(g),
        "Hits": int(g["is_hit"].sum()),
        "Rows in Hit+Miss tabs": len(full),
        "Reviewed": int(full["reviewed"].sum()),
        "% reviewed": _rate(full["reviewed"].sum(), len(full)),
        "Misses": misses,
        "Maybes": int(full["is_maybe"].sum()),
        "Hit rate": _rate(hits, hits + misses),
        "Hits with no corrections %": g["kept_all"].mean(),
        "Hits flagged malfunction": int(normalise_tag(g.loc[g["is_hit"], "is_malfunction"], "is_malfunction").sum()),
        "Hits flagged model comparison": int(normalise_tag(g.loc[g["is_hit"], "model_comparison"], "model_comparison").sum()),
    }
    for field, label in TAG_FIELDS:
        out[f"{label} accuracy"] = g[f"kept_{field}"].mean()
    return pd.Series(out)


def autovalidation_metrics(g: pd.DataFrame) -> pd.Series:
    """How well validation_auto agrees with the analysts. Precision needs Misses, so it is
    only computed on Hit+Miss tabs; recall (Hits it caught) works on every tab."""
    hits = g[g["is_hit"]]
    caught = hits["auto"].isin(["Hit", "Maybe"])
    full = g[g["full_review"] & (g["is_hit"] | g["is_miss"])]
    auto_hit = full[full["auto"] == "Hit"]
    decided = full[full["auto"].isin(["Hit", "Miss"])]
    prefiltered = hits["autoval_status"].str.startswith("prefilter")
    hit_rows = hits[hits["in_llm_file"]]
    out = {
        "Rows": len(g),
        "Auto Hit share": _rate((g["auto"] == "Hit").sum(), len(g)),
        "Auto blank": int((g["auto"] == "").sum()),
        "Analyst Hits": len(hits),
        "Recall (Hits auto caught)": _rate(caught.sum(), len(hits)),
        "Hits auto marked Miss": int((hits["auto"] == "Miss").sum()),
        "...of which by prefilter": int(((hits["auto"] == "Miss") & prefiltered).sum()),
        "Precision (auto Hits that were Hits)": _rate(auto_hit["is_hit"].sum(), len(auto_hit)),
        "Agreement (Hit+Miss tabs)": _rate((decided["auto"] == decided["validation"]).sum(), len(decided)),
        "Auto Miss & unreviewed": int(((g["auto"] == "Miss") & ~g["reviewed"]).sum()),
        "Malfunction auto accuracy (Hits)": _rate(
            (normalise_tag(hit_rows["is_malfunction_auto"], "is_malfunction") == normalise_tag(hit_rows["is_malfunction"], "is_malfunction")).sum(),
            len(hit_rows)),
        "Model comparison auto accuracy (Hits)": _rate(
            (normalise_tag(hit_rows["model_comparison_auto"], "model_comparison") == normalise_tag(hit_rows["model_comparison"], "model_comparison")).sum(),
            len(hit_rows)),
    }
    return pd.Series(out)


def confidence_table(rows: pd.DataFrame) -> pd.DataFrame:
    """Analyst hit rate by what autovalidation said and how confident it was (Hit+Miss tabs only)."""
    full = rows[rows["full_review"] & (rows["is_hit"] | rows["is_miss"])].copy()
    full["Confidence band"] = pd.cut(full["confidence"], [0, 69, 89, 94, 100], labels=["<70", "70-89", "90-94", "95-100"])
    full["Auto label"] = full["auto"].replace("", "(blank)")
    t = full.groupby(["Auto label", "Confidence band"], observed=True).agg(
        Reviewed=("validation", "size"), Hits=("is_hit", "sum"), Misses=("is_miss", "sum"))
    t["Hit rate"] = t["Hits"] / t["Reviewed"]
    return t.reset_index()


def status_table(rows: pd.DataFrame) -> pd.DataFrame:
    """Outcomes by how autovalidation decided each row (llm call vs one of the prefilters)."""
    s = rows.assign(status=rows["autoval_status"].replace("", "(blank)"))
    t = s.groupby("status").agg(
        Rows=("tab", "size"), Reviewed=("reviewed", "sum"), Hits=("is_hit", "sum"), Misses=("is_miss", "sum"))
    t["Hit rate (reviewed, Hit+Miss tabs)"] = s[s["full_review"]].groupby("status").apply(
        lambda g: _rate(g["is_hit"].sum(), (g["is_hit"] | g["is_miss"]).sum()), include_groups=False)
    return t.reset_index().rename(columns={"status": "autoval_status"})


def subcategory_corrections(rows: pd.DataFrame) -> pd.DataFrame:
    changed = rows[rows["kept_feedback_subcategory"] == 0]
    t = (changed.groupby(["llm_feedback_subcategory", "feedback_subcategory"]).size()
         .rename("Hits corrected").reset_index()
         .rename(columns={"llm_feedback_subcategory": "LLM subcategory", "feedback_subcategory": "Analyst subcategory"}))
    totals = rows[rows["is_hit"]].groupby("llm_subcategory").size()
    t["% of LLM subcategory's Hits"] = t["Hits corrected"] / t["LLM subcategory"].map(totals)
    return t.sort_values("Hits corrected", ascending=False)


def corrections_detail(rows: pd.DataFrame) -> pd.DataFrame:
    """One line per corrected field on a Hit, for spot checks."""
    out = []
    for field, label in TAG_FIELDS:
        c = rows[rows[f"kept_{field}"] == 0]
        text = c["translated_text"].where(c["translated_text"] != "", c["post_text"]).str.slice(0, 300)
        out.append(pd.DataFrame({
            "Tab": c["tab"], "id-3": c["id-3"], "Field": label,
            "LLM value": c[f"llm_{field}"], "Analyst value": c[field],
            "LLM subcategory": c["llm_subcategory"], "Text": text, "Sidenote": c["Sidenote"], "URL": c["url"],
        }))
    return pd.concat(out, ignore_index=True).sort_values(["Tab", "Field"])


def data_quality(rows: pd.DataFrame, missing: list[str], empty: list[str], untagged: list[str], week: str) -> pd.DataFrame:
    flags = []

    def flag(severity, tab, issue, count=None, example=""):
        flags.append({"Severity": severity, "Tab": tab, "Issue": issue, "Count": count, "Example": example})

    for field in untagged:
        flag("Info", "All", f"LLM file has the same {field} on every row — not tagged upstream, so left out of tag accuracy")

    for tab in missing:
        flag("Warning", tab, "In vehicle_ids.csv but no tab in the sheet")
    for tab in empty:
        flag("Info", tab, "Tab is empty")

    for tab, g in rows.groupby("tab"):
        mode = g["review_mode"].iloc[0]
        if mode == "Not reviewed":
            flag("Info" if tab in IRREGULAR_TABS else "Warning", tab, "No rows marked Hit or Miss", len(g))
        elif mode == "Hits only" and week >= ALL_TABS_MARK_MISSES_FROM:
            flag("Warning", tab, "Hits marked but no Misses — rejected rows may have been left blank", int(g["is_hit"].sum()))

        invalid = g[g["validation"] == "Invalid"]
        if len(invalid):
            flag("Warning", tab, "Unrecognised value in Validation column", len(invalid),
                 ", ".join(invalid["Validation"].value_counts().index[:3]))

        not_found = g[~g["in_llm_file"]]
        if len(not_found):
            flag("Info", tab, "Rows not found in the week's LLM file (id-3) — excluded from tag accuracy", len(not_found))

        dup = g["id-3"].duplicated(keep=False) & (g["id-3"] != "")
        if dup.any():
            flag("Info", tab, "Duplicate id-3 rows", int(dup.sum()))

        if len(g) >= 50:
            top = g["auto"].value_counts(normalize=True)
            if top.iloc[0] >= 0.99 and top.index[0] != "":
                flag("Warning", tab, f"Autovalidation gave {top.index[0]!r} to every row", len(g))
        blank_auto = (g["auto"] == "").sum()
        if blank_auto > max(5, 0.05 * len(g)):
            flag("Warning", tab, "Autovalidation left rows blank (not run / errored?)", int(blank_auto))

        hits = g[g["is_hit"]]
        bad_sent = hits[~hits["feedback_sentiment"].isin(VALID_SENTIMENTS) | ~hits["overall_sentiment"].isin(VALID_SENTIMENTS)]
        if len(bad_sent):
            flag("Warning", tab, "Hits with an invalid sentiment — excel_upload.py will reject them", len(bad_sent),
                 ", ".join(bad_sent["feedback_sentiment"].unique()[:3]))
        bad_own = hits[~hits["ownership_status"].isin(VALID_OWNERSHIP)]
        if len(bad_own):
            flag("Warning", tab, "Hits with an ownership status excel_upload.py rejects", len(bad_own),
                 ", ".join(bad_own["ownership_status"].unique()[:3]))

    order = {"Warning": 0, "Info": 1}
    df = pd.DataFrame(flags, columns=["Severity", "Tab", "Issue", "Count", "Example"])
    return df.sort_values(["Severity", "Tab"], key=lambda s: s.map(order) if s.name == "Severity" else s)


# ──────────────────────────────────────────────────────────────────────────────
# History
# ──────────────────────────────────────────────────────────────────────────────

def update_history(week: str, overall: pd.Series, by_vehicle: pd.DataFrame, auto_overall: pd.Series) -> pd.DataFrame:
    """One row per week × tab ("ALL" = every tab). Re-running a week replaces its rows."""
    current = pd.concat([
        pd.concat([overall, auto_overall[["Recall (Hits auto caught)", "Precision (auto Hits that were Hits)"]]])
          .to_frame().T.assign(Tab="ALL"),
        by_vehicle.reset_index().rename(columns={"tab": "Tab"}),
    ], ignore_index=True)
    current.insert(0, "Week", week)
    if HISTORY_PATH.exists():
        old = pd.read_csv(HISTORY_PATH, dtype={"Week": str})
        current = pd.concat([old[old["Week"] != week], current], ignore_index=True)
    current = current.infer_objects().sort_values(["Week", "Tab"])
    current.to_csv(HISTORY_PATH, index=False)
    return current


def trend_tables(history: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
    overall = history[history["Tab"] == "ALL"].drop(columns="Tab").set_index("Week")
    tabs = history[history["Tab"] != "ALL"]
    return [
        ("Overall by week", overall.reset_index()),
        ("Hit rate by tab", tabs.pivot(index="Tab", columns="Week", values="Hit rate").reset_index()),
        ("Subcategory accuracy by tab", tabs.pivot(index="Tab", columns="Week", values="Subcategory accuracy").reset_index()),
        ("Rows reviewed by tab", tabs.pivot(index="Tab", columns="Week", values="Reviewed").reset_index()),
    ]


# ──────────────────────────────────────────────────────────────────────────────
# Excel
# ──────────────────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)

DEFINITIONS = [
    ("Hit rate", "Hits ÷ (Hits + Misses). Only from tabs that mark Misses — in Hits-only tabs a blank could be a rejected row."),
    ("% reviewed", "Rows marked Hit, Miss or Maybe ÷ all rows, in tabs that mark Misses. Blank = not reviewed."),
    ("<Field> accuracy", "Share of Hits where the analyst left the LLM's value unchanged (compared with the week's added_facebook_names file via id-3)."),
    ("Hits with no corrections %", "Share of Hits where none of the tagged fields were changed."),
    ("Subcategory metrics", "Grouped by the LLM's original subcategory, so corrected rows count against the tag the LLM gave them."),
    ("Recall", "Analyst Hits that autovalidation marked Hit or Maybe."),
    ("Precision", "Of the rows autovalidation marked Hit and an analyst reviewed, the share that were Hits (Hit+Miss tabs only)."),
]


def _write_table(ws, df: pd.DataFrame, start_row: int, title: str | None = None) -> int:
    """Write df at start_row (1-based) with a styled header; returns the next free row."""
    if title:
        ws.cell(start_row, 1, title).font = TITLE_FONT
        start_row += 1
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start_row, j, str(col))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (float, np.floating)) and np.isnan(v):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            elif isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            c = ws.cell(i, j, v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"  # post text, not a formula
    end_row = start_row + len(df)
    for j, col in enumerate(df.columns, 1):
        name = str(col)
        is_pct = any(k in name for k in ("%", "rate", "accuracy", "Accuracy", "Recall", "Precision", "Agreement", "share"))
        # Week-pivot tables hold percentages under date headers
        if title and ("rate" in title or "accuracy" in title) and j > 1:
            is_pct = True
        if is_pct:
            for i in range(start_row + 1, end_row + 1):
                ws.cell(i, j).number_format = "0.0%"
            if len(df) > 1:
                rng = f"{get_column_letter(j)}{start_row + 1}:{get_column_letter(j)}{end_row}"
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="num", start_value=0, start_color="F8696B",
                    mid_type="num", mid_value=0.5, mid_color="FFEB84",
                    end_type="num", end_value=1, end_color="63BE7B"))
        elif pd.api.types.is_numeric_dtype(df[col]):
            for i in range(start_row + 1, end_row + 1):
                ws.cell(i, j).number_format = "#,##0"
    return end_row + 2


def _finish_sheet(ws, widths: dict[int, int] | None = None, freeze: str | None = "B2", autofilter: bool = True):
    for j, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
        longest = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, longest + 2), 60)
    for j, w in (widths or {}).items():
        ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = freeze
    if autofilter and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def write_excel(path: Path, week: str, sheet_url: str, overall: pd.Series, auto_overall: pd.Series,
                tables: dict[str, pd.DataFrame], trends: list[tuple[str, pd.DataFrame]]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        wb = xw.book
        ws = wb.create_sheet("Summary")
        ws.cell(1, 1, f"Ford tagging accuracy — w/c {week}").font = Font(bold=True, size=15)
        ws.cell(2, 1, "Tagging sheet")
        link = ws.cell(2, 2, sheet_url)
        link.hyperlink, link.font = sheet_url, Font(color="0563C1", underline="single")
        ws.cell(3, 1, f"Generated {datetime.datetime.now():%Y-%m-%d %H:%M}")
        headline = pd.concat([overall, auto_overall.drop(["Rows"])]).reset_index()
        headline.columns = ["Metric", "Value"]
        row = 5
        ws.cell(row, 1, "Headline").font = TITLE_FONT
        for i, (metric, value) in enumerate(headline.itertuples(index=False), row + 1):
            ws.cell(i, 1, metric)
            c = ws.cell(i, 2, None if pd.isna(value) else float(value))
            pct = any(k in metric for k in ("%", "rate", "accuracy", "Recall", "Precision", "Agreement", "share"))
            c.number_format = "0.0%" if pct else "#,##0"
        row += len(headline) + 2
        ws.cell(row, 1, "Definitions").font = TITLE_FONT
        for i, (term, text) in enumerate(DEFINITIONS, row + 1):
            ws.cell(i, 1, term).font = Font(bold=True)
            ws.cell(i, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 110

        for name, df in tables.items():
            ws = wb.create_sheet(name)
            _write_table(ws, df, 1)
            widths = {7: 80} if name == "Corrections Detail" else None
            _finish_sheet(ws, widths=widths)

        ws = wb.create_sheet("Trend")
        row = 1
        for title, df in trends:
            row = _write_table(ws, df, row, title)
        _finish_sheet(ws, freeze=None, autofilter=False)

        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]


# ──────────────────────────────────────────────────────────────────────────────
# Slack
# ──────────────────────────────────────────────────────────────────────────────

def _pct(x) -> str:
    return "n/a" if pd.isna(x) else f"{x:.0%}"


def _delta(now, before, pct=True) -> str:
    if pd.isna(now) or before is None or pd.isna(before):
        return ""
    d = now - before
    if pct:
        return f" ({'+' if d >= 0 else '−'}{abs(d) * 100:.0f}pp)"
    return f" ({'+' if d >= 0 else '−'}{abs(d):,.0f})"


def slack_summary(week: str, sheet_url: str, overall: pd.Series, auto_overall: pd.Series,
                  by_vehicle: pd.DataFrame, by_subcat: pd.DataFrame, flags: pd.DataFrame,
                  history: pd.DataFrame) -> str:
    prev_weeks = sorted(w for w in history["Week"].unique() if w < week)
    prev = history[(history["Week"] == prev_weeks[-1]) & (history["Tab"] == "ALL")].iloc[0] if prev_weeks else None

    def p(metric):
        return None if prev is None else prev.get(metric)

    week_label = datetime.date.fromisoformat(week).strftime("%-d %b %Y")
    lines = [f"*Ford tagging accuracy — w/c {week_label}*  (<{sheet_url}|tagging sheet>)", ""]
    full_hits = overall["Reviewed"] - overall["Misses"] - overall["Maybes"]
    hits_only_tabs = by_vehicle.index[by_vehicle["Rows in Hit+Miss tabs"] == 0]
    extra = overall["Hits"] - full_hits
    lines.append(
        f"• *Reviewed:* {overall['Reviewed']:,.0f} of {overall['Rows in Hit+Miss tabs']:,.0f} rows "
        f"({_pct(overall['% reviewed'])}){_delta(overall['Reviewed'], p('Reviewed'), pct=False)} — "
        f"{full_hits:,.0f} Hits, {overall['Misses']:,.0f} Misses"
        + (f" · plus {extra:,.0f} Hits in tabs that don't mark Misses" if extra and len(hits_only_tabs) else ""))
    lines.append(f"• *Hit rate:* {_pct(overall['Hit rate'])}{_delta(overall['Hit rate'], p('Hit rate'))}")
    lines.append(
        f"• *Tag accuracy on Hits:* subcategory {_pct(overall['Subcategory accuracy'])}"
        f"{_delta(overall['Subcategory accuracy'], p('Subcategory accuracy'))}, "
        f"feedback sentiment {_pct(overall['Feedback sentiment accuracy'])}, "
        f"ownership {_pct(overall['Ownership accuracy'])}, country {_pct(overall['Country accuracy'])} · "
        f"{_pct(overall['Hits with no corrections %'])} of Hits needed no correction")
    lines.append(
        f"• *Autovalidation:* caught {_pct(auto_overall['Recall (Hits auto caught)'])} of Hits "
        f"({auto_overall['Hits auto marked Miss']:,.0f} marked Miss), precision "
        f"{_pct(auto_overall['Precision (auto Hits that were Hits)'])}")

    ranked = by_subcat[by_subcat["Hits"] >= MIN_HITS_TO_RANK].sort_values("Subcategory accuracy")
    if len(ranked):
        worst = ", ".join(f"{s} {_pct(r['Subcategory accuracy'])} (n={r['Hits']:.0f})" for s, r in ranked.head(3).iterrows())
        lines.append(f"• *Least accurate subcategories:* {worst}")
    noisy = by_subcat[by_subcat["Reviewed"] >= MIN_REVIEWED_TO_RANK].sort_values("Hit rate")
    if len(noisy):
        worst = ", ".join(f"{s} {_pct(r['Hit rate'])} (n={r['Reviewed']:.0f})" for s, r in noisy.head(3).iterrows())
        lines.append(f"• *Lowest hit rate subcategories:* {worst}")
    reviewed_tabs = by_vehicle[by_vehicle["Reviewed"] >= MIN_REVIEWED_TO_RANK].sort_values("Hit rate")
    if len(reviewed_tabs):
        worst = ", ".join(f"{t} {_pct(r['Hit rate'])}" for t, r in reviewed_tabs.head(3).iterrows())
        lines.append(f"• *Lowest hit rate vehicles:* {worst}")

    warnings = flags[flags["Severity"] == "Warning"]
    if len(warnings):
        lines += ["", f"*Flags ({len(warnings)})*"]
        for f in warnings.head(6).itertuples():
            count = f" ({f.Count:,.0f})" if pd.notna(f.Count) else ""
            lines.append(f"• {f.Tab}: {f.Issue}{count}")
        if len(warnings) > 6:
            lines.append(f"• …and {len(warnings) - 6} more in the Data Quality tab")
    return "\n".join(lines)


def post_to_slack(text: str, excel_path: Path) -> None:
    from slack_sdk import WebClient

    token = os.getenv("SLACK_BOT_TOKEN")
    if not token:
        print("SLACK_BOT_TOKEN not set — skipping Slack post")
        return
    WebClient(token=token).files_upload_v2(
        channel=SLACK_CHANNEL, file=str(excel_path), title=excel_path.name, initial_comment=text)
    print(f"Posted summary to Slack channel {SLACK_CHANNEL}")


# ──────────────────────────────────────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────────────────────────────────────

def main(week: str | None = None, slack: bool = True) -> Path:
    week = week or last_monday_str()
    print(f"Tagging accuracy for w/c {week}")

    # Sheet first — it fails fast if the week has no link, before the slow xlsx read
    print("Reading the tagging sheet...")
    raw_rows, missing, empty, sheet_url = load_sheet(week)
    print("Loading LLM tags...")
    llm = load_llm_tags(week)
    untagged = untagged_fields(llm)
    rows = prepare_rows(raw_rows, llm, untagged)

    overall = summarise(rows)
    by_vehicle = rows.groupby("tab", sort=False).apply(summarise, include_groups=False)
    by_vehicle.insert(0, "Review mode", rows.groupby("tab", sort=False)["review_mode"].first())
    by_subcat = rows.groupby("llm_subcategory").apply(summarise, include_groups=False).sort_values("Rows", ascending=False)
    by_vehicle_subcat = rows.groupby(["tab", "llm_subcategory"], sort=False).apply(summarise, include_groups=False)
    by_vehicle_subcat = by_vehicle_subcat[(by_vehicle_subcat["Hits"] > 0) | (by_vehicle_subcat["Reviewed"] > 0)]

    # Fields the LLM didn't tag have no accuracy to report
    untagged_cols = [f"{label} accuracy" for f, label in TAG_FIELDS if f in untagged]
    overall = overall.drop(untagged_cols)
    by_vehicle, by_subcat, by_vehicle_subcat = (
        df.drop(columns=untagged_cols) for df in (by_vehicle, by_subcat, by_vehicle_subcat))

    auto_overall = autovalidation_metrics(rows)
    auto_by_vehicle = rows.groupby("tab", sort=False).apply(autovalidation_metrics, include_groups=False)
    flags = data_quality(rows, missing, empty, untagged, week)

    REPORT_DIR.mkdir(exist_ok=True)
    history = update_history(week, overall, by_vehicle.drop(columns="Review mode"), auto_overall)

    excel_path = REPORT_DIR / f"tagging_accuracy_{week}.xlsx"
    tables = {
        "By Vehicle": by_vehicle.reset_index().rename(columns={"tab": "Tab"}),
        "By Subcategory": by_subcat.reset_index().rename(columns={"llm_subcategory": "LLM subcategory"}),
        "Vehicle x Subcategory": by_vehicle_subcat.reset_index().rename(columns={"tab": "Tab", "llm_subcategory": "LLM subcategory"}),
        "Subcategory Corrections": subcategory_corrections(rows),
        "Autovalidation": auto_by_vehicle.reset_index().rename(columns={"tab": "Tab"}),
        "Autoval by Confidence": confidence_table(rows),
        "Autoval by Status": status_table(rows),
        "Data Quality": flags,
        "Corrections Detail": corrections_detail(rows),
    }
    write_excel(excel_path, week, sheet_url, overall, auto_overall, tables, trend_tables(history))
    print(f"Wrote {excel_path}")

    text = slack_summary(week, sheet_url, overall, auto_overall,
                         by_vehicle.drop(columns="Review mode"), by_subcat, flags, history)
    print("\n" + text + "\n")
    if slack:
        post_to_slack(text, excel_path)
    return excel_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--week", help="Monday of the week to report on (YYYY-MM-DD). Defaults to last week.")
    parser.add_argument("--no-slack", action="store_true", help="Write the Excel but don't post to Slack")
    args = parser.parse_args()
    main(args.week, slack=not args.no_slack)
